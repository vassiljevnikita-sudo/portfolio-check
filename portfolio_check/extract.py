"""Pull URLs and text out of Markdown, HTML and .docx files."""

from __future__ import annotations

import re
import zipfile
from pathlib import Path
from urllib.parse import urljoin, urldefrag

FENCE_RE = re.compile(r"^(```|~~~).*?^\1\s*$", re.M | re.S)
MD_LINK_RE = re.compile(r"!?\[[^\]]*\]\(\s*<?([^)\s>]+)>?(?:\s+\"[^\"]*\")?\s*\)")
MD_REF_RE = re.compile(r"^\s*\[[^\]]+\]:\s*<?(\S+?)>?(?:\s+\"[^\"]*\")?\s*$", re.M)
HTML_ATTR_RE = re.compile(r"""\b(?:href|src)\s*=\s*["']([^"']+)["']""", re.I)
BARE_URL_RE = re.compile(r"""(?<![("'<\[=])\bhttps?://[^\s<>"'`)\]]+""")
TEXT_URL_RE = re.compile(
    r"""(?:https?://|www\.)[^\s<>"'`)\]]+|(?<![\w@./-])(?:github\.com|linkedin\.com)/[^\s<>"'`)\]]+""",
    re.I,
)
TRAILING = ".,;:!?*_"


def _clean(url: str) -> str:
    url = url.strip().rstrip(TRAILING)
    if url.lower().startswith("www."):
        url = "https://" + url
    if re.match(r"(?i)^(github|linkedin)\.com/", url):
        url = "https://" + url
    return url


def strip_code_fences(markdown: str) -> str:
    return FENCE_RE.sub("", markdown)


def markdown_links(markdown: str) -> list[str]:
    """All link and image targets in a Markdown document, code blocks excluded."""
    body = strip_code_fences(markdown)
    body = re.sub(r"`[^`\n]*`", "", body)  # inline code
    found: list[str] = []
    for regex in (MD_LINK_RE, MD_REF_RE, HTML_ATTR_RE, BARE_URL_RE):
        found.extend(m.group(1) if regex.groups else m.group(0) for m in regex.finditer(body))
    return _unique(_clean(u) for u in found if u.strip())


def text_urls(text: str) -> list[str]:
    return _unique(_clean(m.group(0)) for m in TEXT_URL_RE.finditer(text))


def _unique(items) -> list[str]:
    seen: dict[str, None] = {}
    for item in items:
        if item and item not in seen:
            seen[item] = None
    return list(seen)


def normalize_for_fetch(url: str) -> str:
    return urldefrag(url)[0]


def resolve_readme_link(link: str, slug: str, readme_path: str = "README.md") -> str:
    """Turn a relative README link into the URL GitHub would show a visitor."""
    if re.match(r"(?i)^[a-z][a-z0-9+.-]*:", link) or link.startswith("//"):
        return link
    if link.startswith("#"):
        return f"https://github.com/{slug}{link}"
    base = f"https://github.com/{slug}/blob/HEAD/{readme_path}"
    return urljoin(base, link)


# --------------------------------------------------------------------- docx

W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
HYPERLINK_FIELD_RE = re.compile(r'HYPERLINK\s+"([^"]+)"')


def docx_text(path: str | Path) -> str:
    """Body paragraphs, table cells (nested too), headers and footers, in reading order."""
    from docx import Document

    doc = Document(str(path))
    lines: list[str] = []

    def walk(container):
        for block in container.iter_inner_content():
            if hasattr(block, "rows"):
                for row in block.rows:
                    seen_cells = set()
                    for cell in row.cells:
                        if id(cell._tc) in seen_cells:
                            continue
                        seen_cells.add(id(cell._tc))
                        walk(cell)
            else:
                lines.append(block.text)

    for section in doc.sections:
        for part in (section.header, section.first_page_header):
            if not part.is_linked_to_previous:
                walk(part)
    walk(doc)
    for section in doc.sections:
        for part in (section.footer, section.first_page_footer):
            if not part.is_linked_to_previous:
                walk(part)
    return "\n".join(lines)


def docx_paragraphs(path: str | Path) -> list[tuple[str, str]]:
    """(style name, text) for every paragraph, including those in table cells."""
    from docx import Document

    doc = Document(str(path))
    out: list[tuple[str, str]] = []

    def walk(container):
        for block in container.iter_inner_content():
            if hasattr(block, "rows"):
                for row in block.rows:
                    for cell in row.cells:
                        walk(cell)
            else:
                style = block.style.name if block.style is not None else ""
                out.append((style, block.text))

    walk(doc)
    return out


def docx_links(path: str | Path) -> list[str]:
    """Hyperlink relationship targets, HYPERLINK field codes and URLs typed in the text."""
    found: list[str] = []
    with zipfile.ZipFile(path) as zf:
        for name in zf.namelist():
            if name.startswith("word/_rels/") and name.endswith(".rels"):
                xml = zf.read(name).decode("utf-8", "replace")
                for m in re.finditer(r"<Relationship\b[^>]*>", xml):
                    tag = m.group(0)
                    if 'TargetMode="External"' in tag and "/hyperlink" in tag:
                        target = re.search(r'Target="([^"]+)"', tag)
                        if target:
                            found.append(_unescape_xml(target.group(1)))
            elif name.startswith("word/") and name.endswith(".xml"):
                xml = zf.read(name).decode("utf-8", "replace")
                found.extend(_unescape_xml(u) for u in HYPERLINK_FIELD_RE.findall(xml))
    found.extend(text_urls(docx_text(path)))
    return _unique(_clean(u) for u in found)


def docx_has_nested_tables(path: str | Path) -> int:
    from docx import Document

    doc = Document(str(path))
    count = 0

    def walk_tables(tables, depth):
        nonlocal count
        for table in tables:
            if depth > 0:
                count += 1
            for row in table.rows:
                for cell in row.cells:
                    walk_tables(cell.tables, depth + 1)

    walk_tables(doc.tables, 0)
    return count


def _unescape_xml(s: str) -> str:
    return s.replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">").replace("&quot;", '"')


def office_zip_text(path: Path) -> str:
    """Raw text from the XML parts of a .docx/.xlsx/.pptx, tags removed."""
    chunks: list[str] = []
    try:
        with zipfile.ZipFile(path) as zf:
            for name in zf.namelist():
                if name.endswith(".xml") or name.endswith(".rels"):
                    xml = zf.read(name).decode("utf-8", "replace")
                    chunks.append(re.sub(r"<[^>]+>", " ", xml))
    except zipfile.BadZipFile:
        return ""
    return "\n".join(chunks)
