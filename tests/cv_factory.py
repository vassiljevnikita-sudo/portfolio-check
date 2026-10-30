"""Build small synthetic CV .docx files for the tests."""

from docx import Document

GOOD_EN = [
    ("h", "Profile"),
    ("p", "Pharmacy student who builds AI tools for quality assurance work. Available up to "
          "15 hours a week from 01.02.2027 to 30.06.2027, full-time from 15.09.2027."),
    ("h", "Projects"),
    ("p", "Questify: task app released for Android and web."),
    ("p", "Complaint Dossier: PowerShell workflow designed along GAMP 5 principles for GMP work."),
    ("p", "Stock selection research: 1.63 billion prediction rows, 2,790 model families. "
          "A prototype was built and presented. Code: github.com/vassiljevnikita-sudo"),
    ("h", "Experience"),
    ("p", "Quality assurance in a pharmaceutical company."),
    ("h", "Education"),
    ("p", "Pharmacy, university studies."),
    ("h", "Skills"),
    ("p", "Python, TypeScript, PowerShell, GitHub."),
    ("h", "Languages"),
    ("p", "German, English."),
]


def make_cv(path, items=GOOD_EN, nested_table=False):
    doc = Document()
    for kind, text in items:
        if kind == "h":
            doc.add_heading(text, level=1)
        else:
            doc.add_paragraph(text)
    if nested_table:
        outer = doc.add_table(rows=1, cols=1)
        outer.cell(0, 0).add_table(rows=1, cols=1)
    doc.save(path)
    return path
