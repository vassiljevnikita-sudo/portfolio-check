import requests

from portfolio_check.checks import links
from portfolio_check.context import RenderResult
from portfolio_check.result import FAIL, PASS

from .conftest import FakeResponse


def _fake_fetch(table):
    def fetch(url, timeout=None):
        entry = table[url]
        if isinstance(entry, Exception):
            raise entry
        status, final = entry if isinstance(entry, tuple) else (entry, url)
        return FakeResponse(status=status, url=final)

    return fetch


def _repo(ctx, tmp_path, body):
    d = tmp_path / "repo"
    d.mkdir()
    (d / "README.md").write_text(body)
    (d / "docs").mkdir()
    (d / "docs" / "shot.png").write_bytes(b"png")
    ctx.local_paths["demo-repo"] = d
    return ctx.config.repos[0]


def test_links_pass_when_all_return_200(ctx, tmp_path):
    repo = _repo(ctx, tmp_path, "# X\n[a](https://a.example/) ![s](docs/shot.png) https://b.example/page\n")
    ctx.fetch = _fake_fetch({"https://a.example/": 200, "https://b.example/page": 200})
    result = links.run(ctx, repo)
    assert result.status == PASS, result.evidence
    assert result.data == {"checked": 2, "broken": 0}


def test_links_fail_on_404_login_redirect_timeout_and_missing_file(ctx, tmp_path):
    repo = _repo(
        ctx,
        tmp_path,
        "# X\n[a](https://a.example/) [p](https://private.example/x) [t](https://slow.example/) [m](docs/gone.md)\n",
    )
    ctx.fetch = _fake_fetch(
        {
            "https://a.example/": 404,
            "https://private.example/x": (200, "https://github.com/login?return_to=x"),
            "https://slow.example/": requests.Timeout(),
        }
    )
    result = links.run(ctx, repo)
    assert result.status == FAIL
    joined = "\n".join(result.evidence)
    assert "404 https://a.example/" in joined
    assert "redirects to login page" in joined
    assert "timeout" in joined
    assert "MISSING local file docs/gone.md" in joined


def test_rate_limit_is_retried():
    calls = []

    class Ctx:
        config = type("C", (), {"link_timeout": 10})()

        def fetch(self, url, timeout=None):
            calls.append(url)
            status = 429 if len(calls) == 1 else 200
            return FakeResponse(status=status, url=url, headers={"Retry-After": "1"})

    ok, line = links.check_url(Ctx(), "https://github.com/x", sleep=lambda s: None)
    assert ok and len(calls) == 2


def test_site_links_come_from_the_rendered_page(ctx):
    site = ctx.config.sites[0]
    ctx.render = lambda s: RenderResult(s.url, s.url, "", "", ["https://demo.example/#top", "mailto:a@b.c", "https://x.example/404"])
    ctx.fetch = _fake_fetch({"https://demo.example/": 200, "https://x.example/404": 404})
    result = links.run(ctx, site)
    assert result.status == FAIL
    assert any(e.startswith("skip mailto:") for e in result.evidence)


def test_cv_links_include_relationships_and_typed_urls(ctx, tmp_path):
    from docx import Document
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = Document()
    p = doc.add_paragraph("Profile: ")
    rid = doc.part.relate_to("https://rel.example/profile", RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "my profile"
    r.append(t)
    h.append(r)
    p._p.append(h)
    doc.add_paragraph("Code: github.com/someone and www.typed.example/page.")
    path = tmp_path / "cv.docx"
    doc.save(path)

    found = links.docx_links(path)
    assert "https://rel.example/profile" in found
    assert "https://github.com/someone" in found
    assert "https://www.typed.example/page" in found


def _scholar_case(ctx, tmp_path, monkeypatch, on_ci):
    repo = _repo(ctx, tmp_path, "# X\n[s](https://scholar.google.com/citations?user=x) [o](https://other.example/)\n")
    ctx.config.link_exceptions = [{"host": "scholar.google.com", "status": [403], "reason": "blocks runners"}]
    monkeypatch.setenv("GITHUB_ACTIONS", "true" if on_ci else "false")
    return repo


def test_link_exception_reports_unverified_on_ci_runner(ctx, tmp_path, monkeypatch):
    repo = _scholar_case(ctx, tmp_path, monkeypatch, on_ci=True)
    ctx.fetch = _fake_fetch({"https://scholar.google.com/citations?user=x": 403, "https://other.example/": 200})
    result = links.run(ctx, repo)
    assert result.status == PASS, result.evidence
    assert result.data == {"checked": 2, "broken": 0, "unverified": 1}
    assert "not verified from CI: blocks runners" in "\n".join(result.evidence)


def test_link_exception_does_not_apply_locally_or_to_other_statuses(ctx, tmp_path, monkeypatch):
    repo = _scholar_case(ctx, tmp_path, monkeypatch, on_ci=False)
    ctx.fetch = _fake_fetch({"https://scholar.google.com/citations?user=x": 403, "https://other.example/": 200})
    assert links.run(ctx, repo).status == FAIL

    monkeypatch.setenv("GITHUB_ACTIONS", "true")
    ctx.fetch = _fake_fetch({"https://scholar.google.com/citations?user=x": 404, "https://other.example/": 200})
    assert links.run(ctx, repo).status == FAIL

    ctx.fetch = _fake_fetch({"https://scholar.google.com/citations?user=x": 403, "https://other.example/": 403})
    result = links.run(ctx, repo)
    assert result.status == FAIL
    assert result.data["broken"] == 1 and result.data["unverified"] == 1
